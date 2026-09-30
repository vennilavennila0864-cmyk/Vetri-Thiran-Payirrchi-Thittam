from io import BytesIO
from docx import Document as DocxDocument
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet

def export_docx(title, sections):
    doc = DocxDocument()
    doc.add_heading(title, 0)
    for section in sections:
        doc.add_heading(section["heading"], level=1)
        for paragraph in section["content"].split("\n"):
            doc.add_paragraph(paragraph)
    output = BytesIO()
    doc.save(output)
    output.seek(0)
    return output

def export_pdf(title, sections):
    output = BytesIO()
    styles = getSampleStyleSheet()
    story = [Paragraph(title, styles["Title"]), Spacer(1, 12)]
    for section in sections:
        story.append(Paragraph(section["heading"], styles["Heading2"]))
        story.append(Paragraph(section["content"].replace("\n", "<br/>"), styles["BodyText"]))
        story.append(Spacer(1, 10))
    SimpleDocTemplate(output).build(story)
    output.seek(0)
    return output

def export_txt(title, sections):
    text = title + "\n\n"
    for section in sections:
        text += section["heading"] + "\n" + section["content"] + "\n\n"
    return BytesIO(text.encode("utf-8"))
